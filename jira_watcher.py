#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Fikstür Talep — Jira İzleyici
Excel log dosyasını izler; yeni satır gelince Jira'da otomatik task açar.
"""

import json
import sys
import threading
import time
from pathlib import Path

import requests
from openpyxl import load_workbook
from PySide6.QtCore import Qt, QTimer, Signal, QObject
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QLineEdit,
    QPushButton, QTextEdit, QFileDialog, QVBoxLayout, QHBoxLayout,
    QFormLayout, QGroupBox, QComboBox, QSpinBox, QMessageBox,
    QFrame
)

# ─── Yapılandırma ────────────────────────────────────────────────

CONFIG_PATH = Path(__file__).parent / "jira_watcher_config.json"

DEFAULT_CONFIG = {
    "jira_url": "",          # örn: https://sirket.atlassian.net  veya  http://jira.sirket.local
    "auth_type": "basic",    # "basic" = Email+Token,  "bearer" = PAT token
    "email": "",
    "api_token": "",
    "project_key": "",       # örn: OFIX
    "issue_type": "Task",
    "excel_path": "",
    "processed_ids": [],
    "poll_seconds": 30,
}

PRIORITY_MAP = {
    "Düşük":  "Low",
    "Normal": "Medium",
    "Yüksek": "High",
    "Kritik": "Critical",
}

# ─── Jira Yardımcıları ───────────────────────────────────────────

def _auth_headers(cfg: dict) -> dict:
    if cfg["auth_type"] == "bearer":
        return {"Authorization": f"Bearer {cfg['api_token']}",
                "Content-Type": "application/json",
                "Accept": "application/json"}
    return {"Content-Type": "application/json", "Accept": "application/json"}


def _auth_tuple(cfg: dict):
    if cfg["auth_type"] == "basic":
        return (cfg["email"], cfg["api_token"])
    return None


def create_jira_issue(cfg: dict, row: dict) -> str:
    """Jira REST API v2 — yeni issue oluştur. Issue anahtarını döner (örn: OFIX-42)."""
    url = cfg["jira_url"].rstrip("/") + "/rest/api/2/issue"

    talep_no = row.get("Talep No", "")
    proje    = row.get("Proje", "")
    lens     = row.get("Lens Doküman No", "")
    urgency  = str(row.get("Aciliyet", "Normal"))

    summary = f"Fikstür Talebi: {talep_no} — {proje} / {lens}"

    desc_lines = [
        f"*Talep No:* {talep_no}",
        f"*Talep Tipi:* {row.get('Talep Tipi', '')}",
        f"*Zaman Damgası:* {row.get('Zaman Damgası', '')}",
        f"*Proje:* {proje}",
        f"*Lens Doküman No:* {lens}",
        f"*Fikstür Adedi:* {row.get('Fikstür Adedi', '')}",
        f"*Kaplama Sistemi:* {row.get('Kaplama Sistemi', '')}",
        f"*Fikstür Malzemesi:* {row.get('Fikstür Malzemesi', '')}",
        f"*Mesh Uygulaması:* {row.get('Mesh Uygulaması', '')}",
        f"*Buss Bar Kalınlığı (mm):* {row.get('Buss Bar Kalınlığı mm', '')}",
        f"*Aciliyet:* {urgency}",
        f"*Not:* {row.get('Not', '')}",
        f"*Talep Klasörü:* {row.get('Talep Klasörü', '')}",
    ]
    description = "\n".join(desc_lines)

    payload = {
        "fields": {
            "project":     {"key": cfg["project_key"]},
            "summary":     summary,
            "description": description,
            "issuetype":   {"name": cfg.get("issue_type", "Task")},
            "priority":    {"name": PRIORITY_MAP.get(urgency, "Medium")},
        }
    }

    resp = requests.post(
        url,
        json=payload,
        auth=_auth_tuple(cfg),
        headers=_auth_headers(cfg),
        timeout=15,
    )
    resp.raise_for_status()
    return resp.json().get("key", "?")


def test_jira_connection(cfg: dict) -> str:
    """Bağlantıyı test et. Başarılıysa kullanıcı adını döner, hata varsa exception fırlatır."""
    url = cfg["jira_url"].rstrip("/") + "/rest/api/2/myself"
    resp = requests.get(
        url,
        auth=_auth_tuple(cfg),
        headers=_auth_headers(cfg),
        timeout=10,
    )
    resp.raise_for_status()
    data = resp.json()
    return data.get("displayName") or data.get("emailAddress") or "OK"


def read_excel_rows(excel_path: str) -> list[dict]:
    """Excel log dosyasını okur; her satırı dict olarak döner."""
    p = Path(excel_path)
    if not p.exists():
        return []
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 2:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        d = {headers[i]: (str(row[i]) if row[i] is not None else "") for i in range(len(headers))}
        if d.get("Talep No", "").strip():
            result.append(d)
    return result


# ─── Sinyal Köprüsü ─────────────────────────────────────────────

class _Bridge(QObject):
    log_signal = Signal(str)   # (mesaj)


# ─── Ana Pencere ─────────────────────────────────────────────────

class JiraWatcher(QMainWindow):

    def __init__(self):
        super().__init__()
        self.setWindowTitle("Fikstür Talep — Jira İzleyici")
        self.setMinimumSize(700, 600)
        self._cfg = self._load_config()
        self._running = False
        self._timer = QTimer(self)
        self._timer.timeout.connect(self._poll)
        self._bridge = _Bridge()
        self._bridge.log_signal.connect(self._append_log)
        self._build_ui()
        self._fill_fields()

    # ── Config ──────────────────────────────────────────────────

    def _load_config(self) -> dict:
        cfg = dict(DEFAULT_CONFIG)
        if CONFIG_PATH.exists():
            try:
                saved = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
                cfg.update(saved)
            except Exception:
                pass
        return cfg

    def _save_config(self):
        CONFIG_PATH.write_text(
            json.dumps(self._cfg, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )

    # ── UI ──────────────────────────────────────────────────────

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setSpacing(10)
        root.setContentsMargins(16, 16, 16, 16)

        # ── Başlık ──
        title = QLabel("Jira İzleyici")
        title.setFont(QFont("Segoe UI", 16, QFont.Bold))
        title.setStyleSheet("color: #1e3a5f;")
        root.addWidget(title)

        sub = QLabel("Excel log dosyasını izler — yeni talep gelince Jira'da otomatik task açar.")
        sub.setStyleSheet("color: #64748b; font-size: 12px;")
        root.addWidget(sub)

        sep = QFrame(); sep.setFrameShape(QFrame.HLine); sep.setStyleSheet("color:#e2e8f0;")
        root.addWidget(sep)

        # ── Jira Ayarları ──
        grp_jira = QGroupBox("Jira Bağlantı Ayarları")
        form = QFormLayout(grp_jira)
        form.setLabelAlignment(Qt.AlignRight)
        form.setSpacing(8)

        self.ed_url    = QLineEdit(); self.ed_url.setPlaceholderText("https://sirket.atlassian.net")
        self.cb_auth   = QComboBox()
        self.cb_auth.addItems(["Basic (E-posta + API Token)", "Bearer Token (PAT)"])
        self.ed_email  = QLineEdit(); self.ed_email.setPlaceholderText("kullanici@sirket.com")
        self.ed_token  = QLineEdit(); self.ed_token.setPlaceholderText("API token veya PAT")
        self.ed_token.setEchoMode(QLineEdit.Password)
        self.ed_proj   = QLineEdit(); self.ed_proj.setPlaceholderText("OFIX")
        self.ed_itype  = QLineEdit(); self.ed_itype.setPlaceholderText("Task")

        self.cb_auth.currentIndexChanged.connect(self._on_auth_change)

        form.addRow("Jira URL:",       self.ed_url)
        form.addRow("Auth Tipi:",      self.cb_auth)
        self.lbl_email = QLabel("E-posta:")
        form.addRow(self.lbl_email,    self.ed_email)
        form.addRow("API Token / PAT:", self.ed_token)
        form.addRow("Proje Anahtarı:", self.ed_proj)
        form.addRow("Issue Tipi:",     self.ed_itype)

        btn_test = QPushButton("Bağlantıyı Test Et")
        btn_test.clicked.connect(self._test_connection)
        form.addRow("", btn_test)

        root.addWidget(grp_jira)

        # ── Excel Dosyası ──
        grp_xl = QGroupBox("Excel Dosyası")
        xl_lay = QHBoxLayout(grp_xl)
        self.ed_excel = QLineEdit()
        self.ed_excel.setPlaceholderText("fikstur_talep_log.xlsx yolunu seçin")
        btn_browse = QPushButton("Gözat…")
        btn_browse.setFixedWidth(90)
        btn_browse.clicked.connect(self._browse_excel)
        xl_lay.addWidget(self.ed_excel)
        xl_lay.addWidget(btn_browse)
        root.addWidget(grp_xl)

        # ── Yoklama Sıklığı + Kontrol ──
        ctrl_lay = QHBoxLayout()
        ctrl_lay.addWidget(QLabel("Kontrol aralığı (sn):"))
        self.sp_interval = QSpinBox()
        self.sp_interval.setRange(10, 3600)
        self.sp_interval.setValue(30)
        self.sp_interval.setFixedWidth(80)
        ctrl_lay.addWidget(self.sp_interval)
        ctrl_lay.addStretch()

        self.btn_start = QPushButton("▶  İzlemeyi Başlat")
        self.btn_start.setFixedHeight(38)
        self.btn_start.setStyleSheet(
            "QPushButton { background:#1e3a5f; color:white; border-radius:8px; font-weight:700; font-size:13px; }"
            "QPushButton:hover { background:#2563eb; }"
            "QPushButton:pressed { background:#1d4ed8; }"
        )
        self.btn_start.clicked.connect(self._toggle)

        self.btn_clear = QPushButton("Geçmişi Temizle")
        self.btn_clear.setFixedHeight(38)
        self.btn_clear.setToolTip("İşlenmiş talep listesini sıfırlar — tüm talepler yeniden gönderilir.")
        self.btn_clear.clicked.connect(self._clear_history)

        ctrl_lay.addWidget(self.btn_clear)
        ctrl_lay.addWidget(self.btn_start)
        root.addLayout(ctrl_lay)

        # ── Log ──
        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        self.log_box.setFont(QFont("Consolas", 10))
        self.log_box.setStyleSheet(
            "background:#0f172a; color:#94a3b8; border-radius:8px; padding:8px;"
        )
        root.addWidget(self.log_box, 1)

    def _fill_fields(self):
        c = self._cfg
        self.ed_url.setText(c.get("jira_url", ""))
        self.cb_auth.setCurrentIndex(0 if c.get("auth_type") == "basic" else 1)
        self.ed_email.setText(c.get("email", ""))
        self.ed_token.setText(c.get("api_token", ""))
        self.ed_proj.setText(c.get("project_key", ""))
        self.ed_itype.setText(c.get("issue_type", "Task") or "Task")
        self.ed_excel.setText(c.get("excel_path", ""))
        self.sp_interval.setValue(int(c.get("poll_seconds", 30)))
        self._on_auth_change()

    def _on_auth_change(self):
        is_basic = self.cb_auth.currentIndex() == 0
        self.lbl_email.setVisible(is_basic)
        self.ed_email.setVisible(is_basic)

    # ── Aksiyonlar ──────────────────────────────────────────────

    def _browse_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Excel Log Dosyası", str(Path.home()),
            "Excel Dosyaları (*.xlsx *.xls)"
        )
        if path:
            self.ed_excel.setText(path)

    def _read_fields_into_cfg(self):
        self._cfg["jira_url"]     = self.ed_url.text().strip()
        self._cfg["auth_type"]    = "basic" if self.cb_auth.currentIndex() == 0 else "bearer"
        self._cfg["email"]        = self.ed_email.text().strip()
        self._cfg["api_token"]    = self.ed_token.text().strip()
        self._cfg["project_key"]  = self.ed_proj.text().strip().upper()
        self._cfg["issue_type"]   = self.ed_itype.text().strip() or "Task"
        self._cfg["excel_path"]   = self.ed_excel.text().strip()
        self._cfg["poll_seconds"] = self.sp_interval.value()

    def _test_connection(self):
        self._read_fields_into_cfg()
        try:
            name = test_jira_connection(self._cfg)
            QMessageBox.information(self, "Bağlantı Başarılı", f"Bağlantı tamam.\nKullanıcı: {name}")
            self._append_log(f"[TEST] Bağlantı başarılı — {name}")
        except Exception as e:
            QMessageBox.critical(self, "Bağlantı Hatası", str(e))
            self._append_log(f"[HATA] Bağlantı testi başarısız: {e}")

    def _toggle(self):
        if self._running:
            self._stop()
        else:
            self._start()

    def _start(self):
        self._read_fields_into_cfg()

        if not self._cfg["jira_url"]:
            QMessageBox.warning(self, "Eksik Bilgi", "Jira URL boş olamaz.")
            return
        if not self._cfg["project_key"]:
            QMessageBox.warning(self, "Eksik Bilgi", "Proje Anahtarı boş olamaz.")
            return
        if not self._cfg["excel_path"]:
            QMessageBox.warning(self, "Eksik Bilgi", "Excel dosyası seçilmedi.")
            return

        self._save_config()
        self._running = True
        self.btn_start.setText("■  Durdur")
        self.btn_start.setStyleSheet(
            "QPushButton { background:#991b1b; color:white; border-radius:8px; font-weight:700; font-size:13px; }"
            "QPushButton:hover { background:#dc2626; }"
        )
        interval_ms = self._cfg["poll_seconds"] * 1000
        self._timer.start(interval_ms)
        self._append_log(f"[BİLGİ] İzleme başladı — her {self._cfg['poll_seconds']} saniyede bir kontrol.")
        # Hemen bir kez çalıştır
        self._poll()

    def _stop(self):
        self._running = False
        self._timer.stop()
        self.btn_start.setText("▶  İzlemeyi Başlat")
        self.btn_start.setStyleSheet(
            "QPushButton { background:#1e3a5f; color:white; border-radius:8px; font-weight:700; font-size:13px; }"
            "QPushButton:hover { background:#2563eb; }"
            "QPushButton:pressed { background:#1d4ed8; }"
        )
        self._append_log("[BİLGİ] İzleme durduruldu.")

    def _clear_history(self):
        reply = QMessageBox.question(
            self, "Geçmişi Temizle",
            "İşlenmiş tüm talep kayıtları silinecek.\n"
            "Bir sonraki kontrolde tüm talepler Jira'ya tekrar gönderilir.\nDevam edilsin mi?",
            QMessageBox.Yes | QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            self._cfg["processed_ids"] = []
            self._save_config()
            self._append_log("[BİLGİ] Geçmiş temizlendi.")

    # ── Yoklama ─────────────────────────────────────────────────

    def _poll(self):
        """Timer callback — arka planda çalışır."""
        threading.Thread(target=self._poll_worker, daemon=True).start()

    def _poll_worker(self):
        try:
            rows = read_excel_rows(self._cfg["excel_path"])
        except Exception as e:
            self._bridge.log_signal.emit(f"[HATA] Excel okunamadı: {e}")
            return

        processed = set(self._cfg.get("processed_ids", []))
        new_rows = [r for r in rows if r.get("Talep No", "") not in processed]

        if not new_rows:
            ts = time.strftime("%H:%M:%S")
            self._bridge.log_signal.emit(f"[{ts}] Yeni talep yok — {len(rows)} satır kontrol edildi.")
            return

        newly_done = []
        for row in new_rows:
            talep_no = row.get("Talep No", "?")
            try:
                key = create_jira_issue(self._cfg, row)
                self._bridge.log_signal.emit(
                    f"[✔] {talep_no} → Jira task açıldı: {key}"
                )
                newly_done.append(talep_no)
            except Exception as e:
                self._bridge.log_signal.emit(
                    f"[HATA] {talep_no} gönderilemedi: {e}"
                )

        if newly_done:
            self._cfg["processed_ids"] = list(processed | set(newly_done))
            self._save_config()

    # ── Log ─────────────────────────────────────────────────────

    def _append_log(self, msg: str):
        self.log_box.append(msg)
        sb = self.log_box.verticalScrollBar()
        sb.setValue(sb.maximum())

    # ── Kapatma ─────────────────────────────────────────────────

    def closeEvent(self, event):
        self._read_fields_into_cfg()
        self._save_config()
        event.accept()


# ─── Giriş Noktası ──────────────────────────────────────────────

def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    w = JiraWatcher()
    w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
